"""Word 解析器：python-docx 提取段落与表格为 Markdown。"""
from __future__ import annotations

from pathlib import Path

from app.parsers.base import DocumentParser, ParseError


class DocxParser(DocumentParser):
    def parse_bytes(self, data: bytes, name: str = "") -> str:
        try:
            from docx import Document
        except ImportError as e:  # pragma: no cover
            raise ParseError("缺少依赖 python-docx，无法解析 Word") from e
        try:
            doc = Document(self._stream(data))
            parts: list[str] = []
            # 段落（跳过空段）
            for para in doc.paragraphs:
                text = para.text.rstrip()
                if text:
                    parts.append(text)
            # 表格 → Markdown 表格
            for table in doc.tables:
                rows: list[str] = []
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    rows.append("| " + " | ".join(cells) + " |")
                if rows:
                    parts.append("\n".join(rows))
            return "\n\n".join(parts)
        except Exception as e:
            raise ParseError(f"Word 解析失败（{name or '文件'}）：{e}") from e
