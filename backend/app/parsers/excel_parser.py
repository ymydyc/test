"""Excel 解析器：.xlsx 用 openpyxl，.xls 用 xlrd，转为 Markdown 表格。"""
from __future__ import annotations

from pathlib import Path

from app.parsers.base import DocumentParser, ParseError


class ExcelParser(DocumentParser):
    def parse_bytes(self, data: bytes, name: str = "") -> str:
        if Path(name or "").suffix.lower() == ".xls":
            return self._parse_xls(data, name)
        return self._parse_xlsx(data, name)

    def _parse_xlsx(self, data: bytes, name: str = "") -> str:
        try:
            from openpyxl import load_workbook
        except ImportError as e:  # pragma: no cover
            raise ParseError("缺少依赖 openpyxl，无法解析 xlsx") from e
        try:
            wb = load_workbook(self._stream(data), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                sheet_lines = [f"## Sheet: {ws.title}"]
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c).strip() for c in row]
                    if any(cells):
                        sheet_lines.append("| " + " | ".join(cells) + " |")
                if len(sheet_lines) > 1:
                    parts.append("\n".join(sheet_lines))
            wb.close()
            return "\n\n".join(parts)
        except Exception as e:
            raise ParseError(f"Excel 解析失败（{name or '文件'}）：{e}") from e

    def _parse_xls(self, data: bytes, name: str = "") -> str:
        try:
            import xlrd
        except ImportError as e:  # pragma: no cover
            raise ParseError("缺少依赖 xlrd，无法解析 xls") from e
        try:
            book = xlrd.open_workbook(file_contents=data)
            parts: list[str] = []
            for sheet in book.sheets():
                sheet_lines = [f"## Sheet: {sheet.name}"]
                for r in range(sheet.nrows):
                    cells = [str(sheet.cell_value(r, c)).strip() for c in range(sheet.ncols)]
                    if any(cells):
                        sheet_lines.append("| " + " | ".join(cells) + " |")
                if len(sheet_lines) > 1:
                    parts.append("\n".join(sheet_lines))
            return "\n\n".join(parts)
        except Exception as e:
            raise ParseError(f"Excel 解析失败（{name or '文件'}）：{e}") from e
